"""PowerPoint deck built from the result tables and figures."""

import logging
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from kairos import config
from kairos.reporting.findings import findings
from kairos.reporting.results import load_results

log = logging.getLogger(__name__)

BG, PANEL, INK, MUTED, ACCENT = (
    RGBColor(0x0B, 0x0F, 0x17),
    RGBColor(0x16, 0x1C, 0x28),
    RGBColor(0xF2, 0xF4, 0xF8),
    RGBColor(0x9A, 0xA4, 0xB5),
    RGBColor(0x3D, 0xD6, 0xC6),
)
W, H = Inches(13.333), Inches(7.5)
BULLET_X, BULLET_W = Inches(7.6), Inches(5.2)


def _text(
    slide: Any,
    text: str | list[str],
    x: Any,
    y: Any,
    w: Any,
    h: Any,
    size: int = 18,
    color: Any = INK,
    bold: bool = False,
) -> Any:
    """Add a text box."""
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = (
            Pt(size),
            bold,
            color,
            "Calibri",
        )
        p.space_after = Pt(8)
    return box


def _slide(prs: Any, title: str) -> Any:
    """Add a blank slide with a title."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(1, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _text(
        s, title, Inches(0.85), Inches(0.45), Inches(11.5), Inches(0.8), 30, INK, True
    )
    return s


def _image(slide: Any, path: Path, x: Any, y: Any, w: Any) -> None:
    """Add a picture on a white panel when the file exists."""
    if path.exists():
        panel = slide.shapes.add_shape(
            1, x - Inches(0.1), y - Inches(0.1), w + Inches(0.2), Inches(4.9)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        panel.line.fill.background()
        slide.shapes.add_picture(str(path), x, y, width=w)


def _table(
    slide: Any, df: pd.DataFrame, x: Any, y: Any, w: Any, size: int = 12
) -> None:
    """Add a formatted table."""
    rows, cols = len(df) + 1, len(df.columns) + 1
    tbl = slide.shapes.add_table(rows, cols, x, y, w, Inches(0.4) * rows).table
    heads = [""] + list(df.columns)
    for j, h in enumerate(heads):
        tbl.cell(0, j).text = h
    for i, (name, row) in enumerate(df.iterrows(), start=1):
        tbl.cell(i, 0).text = str(name)
        for j, v in enumerate(row, start=1):
            tbl.cell(i, j).text = f"{v:.2f}" if abs(v) >= 0.05 or v == 0 else f"{v:.3f}"
    for i in range(rows):
        for j in range(cols):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else PANEL
            for p in cell.text_frame.paragraphs:
                p.font.size, p.font.name = Pt(size), "Calibri"
                p.font.color.rgb = BG if i == 0 else INK
                p.font.bold = i == 0


def _title_slide(prs: Any) -> None:
    """Add the opening slide."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    _text(
        s,
        "Kairos: Regime-Aware Portfolio Intelligence",
        Inches(0.9),
        Inches(2.4),
        Inches(11.5),
        Inches(1.6),
        40,
        INK,
        True,
    )


def _bullets(
    slide: Any,
    lines: list[str],
    x: Any = BULLET_X,
    w: Any = BULLET_W,
    size: int = 18,
) -> None:
    """Add a bullet list."""
    _text(slide, lines, x, Inches(1.8), w, Inches(4.5), size)


def _diagnosis_lines(r: dict[str, Any]) -> list[str]:
    """Bullets for the diagnosis slide."""
    lines = []
    link = r["regime_link"]
    spy = link[link["Asset"] == "SPY"].set_index("Window")
    if {"initial_window", "after"} <= set(spy.index):
        lines.append(
            "SPY stress-to-return correlation: "
            f"{spy.at['initial_window', 'CorrStressNextRet']:+.2f} in the initial window, "
            f"{spy.at['after', 'CorrStressNextRet']:+.2f} afterwards"
        )
    summ = r["ic_summary"]
    for name, row in summ.iterrows():
        lines.append(f"{name.replace('_', ' ')}: mean IC {row['MeanIC']:+.3f}")
    return lines


def build(paths: config.Paths = config.PATHS) -> Any:
    """Assemble the presentation."""
    r = load_results(paths)
    f = findings(r)
    sm = r["summary"]

    def fig(n: str) -> Path:
        """Path of a generated figure."""
        return paths.figures / n

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    _title_slide(prs)

    s = _slide(prs, "Research question")
    _text(
        s,
        [
            "Does regime-aware machine learning allocation improve risk-adjusted returns over classical methods?",
            "Under what market conditions does any advantage hold or break down?",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(3),
        24,
    )

    s = _slide(prs, "Data and universe")
    _text(
        s,
        [
            "Nine assets: SPY, EFA, TLT, IEF, HYG, GLD, DBC, BTC-USD, ETH-USD",
            "Daily adjusted closes from Yahoo Finance, VIX for regimes, 13-week T-bill as the risk-free rate",
            "No simulated or reconstructed series. Panel validated before use",
            "Universe expands as each asset builds a usable track record",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(4),
        22,
    )

    s = _slide(prs, "Regime detection")
    _image(s, fig("03_regime_validation.png"), Inches(0.9), Inches(1.6), Inches(6.3))
    _bullets(
        s,
        [
            "Walk-forward HMM, no look-ahead",
            f"Stress windows flagged: {f['hmm_stress']:.0%}",
            f"Calm controls flagged: {f['hmm_calm']:.0%}",
            f"Volatility rule: {f['vol_stress']:.0%} and {f['vol_calm']:.0%}",
        ],
        size=20,
    )

    s = _slide(prs, "Forecasting skill")
    _image(s, fig("11_forecast_skill.png"), Inches(0.6), Inches(1.6), Inches(8.2))
    _bullets(
        s,
        [
            f"Mean stacked IC: {f['mean_ic']:.3f}",
            f"Significantly negative IC: {f['sig_neg']} of {f['n_assets']} assets",
            f"Significantly positive IC: {f['sig_pos']} of {f['n_assets']} assets",
            f"R2 below the historical mean on {f['r2_neg_assets']} of {f['n_assets']} assets",
            f"Hit rate below always-up baseline: {f['hit_below_base']} of {f['n_assets']}",
        ],
        x=Inches(9.1),
        w=Inches(3.9),
        size=16,
    )

    s = _slide(prs, "Diagnosing the negative IC")
    _image(s, fig("12_ic_diagnostics.png"), Inches(0.6), Inches(1.6), Inches(7.6))
    _bullets(s, _diagnosis_lines(r), x=Inches(8.6), w=Inches(4.3), size=15)

    skill = r["skill_period"]
    lr = r["lstm_repair"]

    def r2(name: str, period: str) -> float:
        """Mean R2 for one forecast and period."""
        row = skill[(skill["Forecast"] == name) & (skill["Period"] == period)]
        return float(row["MeanR2OOS"].iloc[0]) if len(row) else float("nan")

    s = _slide(prs, "Repairing the forecasts")
    _image(s, fig("15_forecast_repair.png"), Inches(0.5), Inches(1.6), Inches(8.2))
    _bullets(
        s,
        [
            f"LSTM dispersion {lr['Disp_raw'].mean():.1f}x to {lr['Disp_fixed'].mean():.2f}x realised returns",
            f"LSTM mean R2 {lr['R2OOS_raw'].mean():.1f} to {lr['R2OOS_fixed'].mean():.2f}",
            f"Calibrated forecast R2: {r2(config.VIEW_COLUMN, 'development'):.3f} before, {r2(config.VIEW_COLUMN, 'holdout'):.3f} after {config.HOLDOUT_START[:4]}",
            (
                "No forecast beats the historical mean in either period"
                if not ((skill["Period"] != "full") & (skill["MeanR2OOS"] > 0)).any()
                else "Some forecasts beat the historical mean in one period"
            ),
        ],
        x=Inches(9.0),
        w=Inches(4.0),
        size=15,
    )

    s = _slide(prs, "Portfolio performance")
    cols = ["AnnReturn", "AnnVol", "Sharpe", "MaxDrawdown"]
    _table(s, sm[cols], Inches(0.9), Inches(1.5), Inches(11.5), 12)

    s = _slide(prs, "Cumulative returns")
    _image(s, fig("07_cumulative_returns.png"), Inches(2.2), Inches(1.6), Inches(8.6))

    s = _slide(prs, "Statistical significance")
    _image(s, fig("09_sharpe_significance.png"), Inches(0.9), Inches(1.6), Inches(6.3))
    _bullets(
        s,
        [f"Regime ablation, Sharpe difference: {f['ablation']}", f["stress_cvar"]],
        size=16,
    )

    s = _slide(prs, "Overlay ablation")
    _image(s, fig("14_overlay_ablation.png"), Inches(0.6), Inches(1.6), Inches(8.2))
    lines = []
    if {"Regime_60_40", "Static_60_40"} <= set(sm.index):
        lines += [
            f"Same stress rule on 60/40: Sharpe {sm.at['Regime_60_40', 'Sharpe']:.2f} against {sm.at['Static_60_40', 'Sharpe']:.2f}",
            f"Max drawdown {sm.at['Regime_60_40', 'MaxDrawdown']:.1%} against {sm.at['Static_60_40', 'MaxDrawdown']:.1%}",
        ]
    if config.NAIVE in sm.index:
        lines.append(
            f"Naive historical-mean views: Sharpe {sm.at[config.NAIVE, 'Sharpe']:.2f}"
        )
    _bullets(s, lines, x=Inches(9.1), w=Inches(3.9), size=16)

    w = r["window"]
    s = _slide(prs, "Other evaluation windows")
    _image(s, fig("16_window_robustness.png"), Inches(0.5), Inches(1.6), Inches(8.2))
    wl = []
    if {"Static_60_40", "Regime_60_40"} <= set(w.index):
        wl = [
            f"Window starts {r['window_meta']['first_decision'].iloc[0]}, includes 2008",
            f"Max drawdown static 60/40: {w.at['Static_60_40', 'MaxDD_crisis_window']:.1%}",
            f"Max drawdown Regime_60_40: {w.at['Regime_60_40', 'MaxDD_crisis_window']:.1%}",
            f"BL_ML_Regime Sharpe: {w.at['BL_ML_Regime', 'Sharpe_crisis_window']:.2f}",
        ]
    _bullets(s, wl, x=Inches(9.0), w=Inches(4.0), size=15)

    s = _slide(prs, "Performance by regime")
    _image(s, fig("10_regime_performance.png"), Inches(2.2), Inches(1.6), Inches(8.6))

    s = _slide(prs, "Parameter sensitivity")
    _image(s, fig("13_sensitivity.png"), Inches(1.6), Inches(1.5), Inches(8.4))
    t = r["sensitivity"]
    _bullets(
        s,
        [
            f"Runs above 60/40 Sharpe: {int((t['Sharpe'] > t['Sharpe_60_40']).sum())} of {len(t)}",
            f"Sharpe range: {t['Sharpe'].min():.2f} to {t['Sharpe'].max():.2f}",
        ],
        x=Inches(10.3),
        w=Inches(2.8),
        size=14,
    )

    sig = r["significance"]
    overlay_wins = any(
        (row["strategy"] in config.KAIROS and row["benchmark"] in config.OVERLAYS)
        and row["diff"] > 0
        and row["p_value"] < 0.05
        for _, row in sig.iterrows()
    )
    overlay_line = (
        "Forecast-driven portfolios beat a mechanical stress overlay at the 5% level."
        if overlay_wins
        else "Forecasts add no measurable value beyond a mechanical stress overlay; the regime layer mainly lowers drawdowns."
    )
    pr = r["prior"][r["prior"]["Variant"] == "BL_ML_Regime"]
    s = _slide(prs, "Prior dependence")
    _image(s, fig("17_prior_comparison.png"), Inches(0.5), Inches(1.6), Inches(8.2))
    _bullets(
        s,
        [
            f"{x.Prior}: Sharpe gain over its prior portfolio {x.Diff:+.2f} (p = {x.p_value:.2f})"
            for x in pr.itertuples()
        ],
        x=Inches(9.0),
        w=Inches(4.0),
        size=14,
    )

    hs = r["holdout_status"].iloc[0]
    s = _slide(prs, "Conclusions")
    _text(
        s,
        [
            f["answer"],
            overlay_line,
            f"Specification frozen at {hs['freeze_date']}; {int(hs['months_since_freeze'])} new months evaluated so far.",
            "Limitations: about 15 years of monthly data, no observed market-capitalisation prior, repairs designed after seeing full-sample results.",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(4.8),
        18,
    )
    return prs


def run(paths: config.Paths = config.PATHS) -> Path:
    """Deck stage."""
    out = paths.reports / "Kairos_Deck.pptx"
    paths.reports.mkdir(parents=True, exist_ok=True)
    build(paths).save(out)
    log.info("deck written to %s", out)
    return out
