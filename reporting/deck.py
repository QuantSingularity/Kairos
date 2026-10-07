from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import config
from reporting.analysis import findings, load_results

BG, PANEL, INK, MUTED, ACCENT = (
    RGBColor(0x0B, 0x0F, 0x17),
    RGBColor(0x16, 0x1C, 0x28),
    RGBColor(0xF2, 0xF4, 0xF8),
    RGBColor(0x9A, 0xA4, 0xB5),
    RGBColor(0x3D, 0xD6, 0xC6),
)
W, H = Inches(13.333), Inches(7.5)


def _text(slide, text, x, y, w, h, size=18, color=INK, bold=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = (
            Pt(size),
            bold,
            color,
            "Calibri",
        )
        p.space_after = Pt(8)
    return box


def _slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(1, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _text(
        s, title, Inches(0.85), Inches(0.45), Inches(11.5), Inches(0.8), 30, INK, True
    )
    return s


def _image(slide, path, x, y, w):
    if path.exists():
        panel = slide.shapes.add_shape(
            1, x - Inches(0.1), y - Inches(0.1), w + Inches(0.2), Inches(4.9)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        panel.line.fill.background()
        slide.shapes.add_picture(str(path), x, y, width=w)


def _table(slide, df, x, y, w, size=12):
    rows, cols = len(df) + 1, len(df.columns) + 1
    tbl = slide.shapes.add_table(rows, cols, x, y, w, Inches(0.4) * rows).table
    heads = [""] + list(df.columns)
    for j, h in enumerate(heads):
        tbl.cell(0, j).text = h
    for i, (name, row) in enumerate(df.iterrows(), start=1):
        tbl.cell(i, 0).text = str(name)
        for j, v in enumerate(row, start=1):
            tbl.cell(i, j).text = f"{v:.2f}" if abs(v) >= 0.05 or v == 0 else f"{v:.3f}"
    for i in range(rows):
        for j in range(cols):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else PANEL
            for p in cell.text_frame.paragraphs:
                p.font.size, p.font.name = Pt(size), "Calibri"
                p.font.color.rgb = BG if i == 0 else INK
                p.font.bold = i == 0


def build(paths=config.PATHS):
    r = load_results(paths)
    f = findings(r)
    fig = lambda n: paths.figures / n
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    _text(
        s,
        "Kairos: Regime-Aware Portfolio Intelligence",
        Inches(0.9),
        Inches(2.4),
        Inches(11.5),
        Inches(1.6),
        40,
        INK,
        True,
    )

    s = _slide(prs, "Research question")
    _text(
        s,
        [
            "Does regime-aware machine learning allocation improve risk-adjusted returns over classical methods?",
            "Under what market conditions does any advantage hold or break down?",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(3),
        24,
    )

    s = _slide(prs, "Data and universe")
    _text(
        s,
        [
            "Nine assets: SPY, EFA, TLT, IEF, HYG, GLD, DBC, BTC-USD, ETH-USD",
            "Daily adjusted closes from Yahoo Finance, VIX for regimes, 13-week T-bill as the risk-free rate",
            "No simulated or reconstructed series. Panel validated before use",
            "Universe expands as each asset builds a usable track record",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(4),
        22,
    )

    s = _slide(prs, "Regime detection")
    _image(s, fig("03_regime_validation.png"), Inches(0.9), Inches(1.6), Inches(6.3))
    _text(
        s,
        [
            f"Walk-forward HMM, no look-ahead",
            f"Stress windows flagged: {f['hmm_stress']:.0%}",
            f"Calm controls flagged: {f['hmm_calm']:.0%}",
            f"Volatility rule: {f['vol_stress']:.0%} and {f['vol_calm']:.0%}",
        ],
        Inches(7.6),
        Inches(1.8),
        Inches(5.2),
        Inches(4),
        20,
    )

    s = _slide(prs, "Forecasting skill")
    _image(s, fig("04_model_ic.png"), Inches(0.9), Inches(1.6), Inches(7.2))
    _text(
        s,
        [
            f"Mean stacked IC: {f['mean_ic']:.3f}",
            f"Significant at 5%: {f['sig_assets']} of {f['n_assets']} assets",
            f"Ensemble beats mean single model: {f['ens_beats']} of {f['n_assets']}",
        ],
        Inches(8.5),
        Inches(1.8),
        Inches(4.4),
        Inches(4),
        20,
    )

    s = _slide(prs, "Portfolio performance")
    cols = ["AnnReturn", "AnnVol", "Sharpe", "MaxDrawdown"]
    _table(s, r["summary"][cols], Inches(0.9), Inches(1.7), Inches(11.5), 13)

    s = _slide(prs, "Cumulative returns")
    _image(s, fig("07_cumulative_returns.png"), Inches(2.2), Inches(1.6), Inches(8.6))

    s = _slide(prs, "Statistical significance")
    _image(s, fig("09_sharpe_significance.png"), Inches(0.9), Inches(1.6), Inches(6.3))
    _text(
        s,
        [f"Regime ablation: {f['ablation']}", f["stress_cvar"]],
        Inches(7.6),
        Inches(1.8),
        Inches(5.2),
        Inches(4),
        18,
    )

    s = _slide(prs, "Performance by regime")
    _image(s, fig("10_regime_performance.png"), Inches(2.2), Inches(1.6), Inches(8.6))

    s = _slide(prs, "Conclusions")
    _text(
        s,
        [
            f["answer"],
            "Limitations: about 15 years of monthly data, approximate market-weight prior, 2008 outside the backtest window.",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(4.5),
        20,
    )
    return prs


def run(paths=config.PATHS):
    out = paths.reports / "Kairos_Deck.pptx"
    paths.reports.mkdir(parents=True, exist_ok=True)
    build(paths).save(out)
    return out
